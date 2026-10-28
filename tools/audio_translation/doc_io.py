"""Read/write repository article files (.txt/.md/.docx). Local-only."""
from __future__ import annotations
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parents[2]
SEARCH_DIRS = [REPO_ROOT / "docs-input", REPO_ROOT / "docs"]
ALLOWED_EXTS = {".txt", ".md", ".docx"}


def list_articles(limit: int = 200) -> list[str]:
    found: list[str] = []
    for base in SEARCH_DIRS:
        if not base.exists():
            continue
        for p in sorted(base.rglob("*")):
            if p.is_file() and p.suffix.lower() in ALLOWED_EXTS and "~$" not in p.name:
                try:
                    found.append(str(p.relative_to(REPO_ROOT)))
                except ValueError:
                    found.append(str(p))
                if len(found) >= limit:
                    return found
    return found


def read_article(rel_or_abs: str) -> str:
    p = Path(rel_or_abs)
    if not p.is_absolute():
        p = REPO_ROOT / rel_or_abs
    if p.suffix.lower() == ".docx":
        from docx import Document
        return "\n\n".join(par.text for par in Document(str(p)).paragraphs).strip()
    return p.read_text(encoding="utf-8")


def save_article(rel_or_abs: str, arabic_text: str, mode: str = "append") -> str:
    """mode: append | overwrite. .docx appends paragraphs; .txt/.md appends text."""
    p = Path(rel_or_abs)
    if not p.is_absolute():
        p = REPO_ROOT / rel_or_abs
    p.parent.mkdir(parents=True, exist_ok=True)
    text = arabic_text.strip() + "\n"
    if p.suffix.lower() == ".docx":
        from docx import Document
        doc = Document(str(p)) if p.exists() else Document()
        if mode == "overwrite" and p.exists():
            doc = Document()
        for para in text.split("\n\n"):
            if para.strip():
                doc.add_paragraph(para.strip())
        doc.save(str(p))
    else:
        if mode == "append" and p.exists():
            prev = p.read_text(encoding="utf-8")
            p.write_text(prev.rstrip() + "\n\n" + text, encoding="utf-8")
        else:
            p.write_text(text, encoding="utf-8")
    return str(p)
